import json
import os
from datetime import datetime, timedelta
import traceback
import difflib

def get_sender_address(item):
    try:
        sender_name = getattr(item, 'SenderName', 'Unknown')
        sender_email = getattr(item, 'SenderEmailAddress', '')
        sender_type = getattr(item, 'SenderEmailType', '')
        
        if sender_type == "EX":
            try:
                sender_email = item.Sender.GetExchangeUser().PrimarySmtpAddress
            except:
                pass
                
        if sender_email:
            return f"{sender_name} <{sender_email}>"
        return sender_name
    except:
        return "Unknown"

def get_receivers_address(item):
    try:
        receivers = []
        for r in item.Recipients:
            name = r.Name
            addr = getattr(r, 'Address', '')
            if addr.lower().startswith('/o='):
                try:
                    addr = r.AddressEntry.GetExchangeUser().PrimarySmtpAddress
                except:
                    pass
            if addr:
                receivers.append(f"{name} <{addr}>")
            else:
                receivers.append(name)
        if receivers:
            return "; ".join(receivers)
        return getattr(item, 'To', 'Unknown')
    except:
        return getattr(item, 'To', 'Unknown')

def is_employee(email_str, exclude_list):
    if not exclude_list or not email_str:
        return False
    email_str_lower = email_str.lower()
    for ex in exclude_list:
        if ex.lower() in email_str_lower:
            return True
    return False

def clean_subject(subject):
    s = subject.lower()
    prefixes = ['re:', 'fw:', 'fwd:', '回复:', '答复:', '转发:']
    for p in prefixes:
        s = s.replace(p, '')
    return s.strip()

def process_folder(folder, cutoff_date, max_emails, sent_folder_id, account_smtp_address, account_name, current_count=0):
    """递归处理文件夹，返回提取的邮件列表 (不写入 Excel)"""
    emails = []
    try:
        items = folder.Items
        items.Sort("[ReceivedTime]", True)
        
        for item in items:
            if getattr(item, 'Class', 0) == 43: # MailItem
                received_time = item.ReceivedTime
                if hasattr(received_time, 'replace'):
                    dt = received_time.replace(tzinfo=None)
                else:
                    dt = datetime(received_time.year, received_time.month, received_time.day, 
                                  received_time.hour, received_time.minute, received_time.second)
                
                if dt < cutoff_date:
                    break
                    
                sender_str = get_sender_address(item)
                receiver_str = get_receivers_address(item)
                subject = getattr(item, 'Subject', 'No Subject')
                mail_time_str = dt.strftime('%Y-%m-%d %H:%M:%S')
                log_time_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                
                is_sent = False
                if getattr(folder, 'EntryID', '') == sent_folder_id:
                    is_sent = True
                elif account_smtp_address and account_smtp_address in sender_str:
                    is_sent = True
                elif account_name in sender_str:
                    is_sent = True
                    
                direction = "发送" if is_sent else "接收"
                folder_name = getattr(folder, 'Name', 'Unknown')
                
                email_dict = {
                    "log_time": log_time_str,
                    "mail_time": mail_time_str,
                    "direction": direction,
                    "sender": sender_str,
                    "receiver": receiver_str,
                    "subject": subject,
                    "folder": folder_name,
                    "reply_time": ""
                }
                emails.append(email_dict)
                
                current_count += 1
                if current_count >= max_emails:
                    return emails, current_count
    except Exception as e:
        pass
        
    try:
        for subfolder in folder.Folders:
            if current_count >= max_emails:
                break
            sub_emails, current_count = process_folder(subfolder, cutoff_date, max_emails, sent_folder_id, account_smtp_address, account_name, current_count)
            emails.extend(sub_emails)
    except:
        pass
        
    return emails, current_count

def show_alert_dialog(alert_emails, all_emails, excel_file, lang="zh"):
    try:
        from PySide6.QtWidgets import QApplication, QDialog, QVBoxLayout, QHBoxLayout, QLabel, QCheckBox, QPushButton, QScrollArea, QWidget, QFrame
        from PySide6.QtCore import Qt
    except ImportError:
        print("无法加载 UI 库，跳过弹窗。")
        return

    import sys
    app = QApplication.instance()
    if not app:
        app = QApplication(sys.argv)
        
    dialog = QDialog()
    tr_window_title = "未回复邮件提醒" if lang == "zh" else "Unanswered Emails Alert"
    dialog.setWindowTitle(tr_window_title)
    dialog.resize(800, 500)
    layout = QVBoxLayout(dialog)
    
    tr_title_text = f"有 {len(alert_emails)} 封客户邮件已超时未回复，请确认：" if lang == "zh" else f"You have {len(alert_emails)} unanswered customer emails. Please confirm:"
    lbl_title = QLabel(tr_title_text)
    lbl_title.setStyleSheet("font-weight: bold; font-size: 16px; margin-bottom: 5px;")
    layout.addWidget(lbl_title)
    
    tr_sub_text = "如果某些邮件不需要回复，请在右侧勾选开关，确认后系统将不再提醒。" if lang == "zh" else "Check the box on the right if a reply is not needed. You won't be reminded again."
    lbl_sub = QLabel(tr_sub_text)
    lbl_sub.setStyleSheet("color: #666; margin-bottom: 10px;")
    layout.addWidget(lbl_sub)
    
    scroll = QScrollArea()
    scroll.setWidgetResizable(True)
    scroll_widget = QWidget()
    scroll_layout = QVBoxLayout(scroll_widget)
    scroll_layout.setSpacing(10)
    
    class CopyLabel(QLabel):
        _active_toast = None

        def __init__(self, prefix_html, display_text, copy_text):
            super().__init__()
            self.copy_text = copy_text
            self.setText(f"{prefix_html}{display_text}")
            self.setTextFormat(Qt.TextFormat.RichText)
            self.setCursor(Qt.CursorShape.PointingHandCursor)
            self.setWordWrap(True)
            self.setStyleSheet("QLabel { border: none; background-color: transparent; }")
            tr_tooltip = "点击复制" if lang == "zh" else "Click to copy"
            self.setToolTip(tr_tooltip)
            self.press_pos = None

        def mousePressEvent(self, event):
            if event.button() == Qt.MouseButton.LeftButton:
                QApplication.clipboard().setText(self.copy_text)
                self.press_pos = event.globalPosition().toPoint()
                
                # 清理之前的提示框，防止重叠
                if CopyLabel._active_toast:
                    CopyLabel._active_toast.hide()
                    CopyLabel._active_toast.deleteLater()
                    
                # 创建一个完全受我们自己生命周期控制的自定义提示框，彻底抛弃 QToolTip 防止闪烁
                tr_copied = "已复制" if lang == "zh" else "Copied"
                toast = QLabel(f"{tr_copied}: {self.copy_text}")
                toast.setWindowFlags(Qt.WindowType.ToolTip | Qt.WindowType.FramelessWindowHint)
                toast.setAttribute(Qt.WidgetAttribute.WA_ShowWithoutActivating)
                toast.setStyleSheet("background-color: #333; color: white; border-radius: 4px; padding: 5px 10px; font-weight: bold;")
                toast.move(self.press_pos.x() + 15, self.press_pos.y() + 15)
                toast.show()
                CopyLabel._active_toast = toast

            super().mousePressEvent(event)

        def mouseReleaseEvent(self, event):
            if event.button() == Qt.MouseButton.LeftButton:
                if self.press_pos and CopyLabel._active_toast:
                    from PySide6.QtCore import QTimer
                    toast_ref = CopyLabel._active_toast
                    # 鼠标抬起后，启动 500ms 的精确倒计时
                    QTimer.singleShot(500, lambda: (toast_ref.hide(), toast_ref.deleteLater()))
                    self.press_pos = None
                    CopyLabel._active_toast = None
            super().mouseReleaseEvent(event)

    import re
    checkboxes = []
    
    tr_time_prefix = "<b>时间:</b> " if lang == "zh" else "<b>Time:</b> "
    tr_sender_prefix = "<b>发件人:</b> " if lang == "zh" else "<b>Sender:</b> "
    tr_subject_prefix = "<b>标题:</b> " if lang == "zh" else "<b>Subject:</b> "
    tr_no_need = "不需要回复" if lang == "zh" else "No Need Reply"
    tr_confirm = "确认并关闭" if lang == "zh" else "Confirm & Close"
    
    for e in alert_emails:
        frame = QFrame()
        frame.setStyleSheet("QFrame { border: 1px solid #ddd; border-radius: 5px; background-color: #fafafa; }")
        row_layout = QHBoxLayout(frame)
        row_layout.setContentsMargins(10, 10, 10, 10)
        
        info_vbox = QVBoxLayout()
        info_vbox.setSpacing(5)
        
        top_hbox = QHBoxLayout()
        lbl_time = QLabel(f"<span style='color:#0055a4;'>{tr_time_prefix}{e['mail_time']}</span>")
        lbl_time.setStyleSheet("border: none; background-color: transparent;")
        
        import html
        raw_sender = e['sender']
        match = re.search(r'<([^>]+)>', raw_sender)
        pure_email = match.group(1).strip() if match else raw_sender.strip()
        
        raw_sender_escaped = html.escape(raw_sender)
        subject_escaped = html.escape(e['subject'])
        
        lbl_sender = CopyLabel(
            prefix_html=f"<span style='color:#d35400;'>{tr_sender_prefix}</span>",
            display_text=f"<span>{raw_sender_escaped}</span>",
            copy_text=pure_email
        )
        
        top_hbox.addWidget(lbl_time)
        top_hbox.addSpacing(20)
        top_hbox.addWidget(lbl_sender, stretch=1)
        
        info_vbox.addLayout(top_hbox)
        
        lbl_subject = CopyLabel(
            prefix_html=tr_subject_prefix,
            display_text=f"<span>{subject_escaped}</span>",
            copy_text=e['subject']
        )
        info_vbox.addWidget(lbl_subject)
        
        chk = QCheckBox(tr_no_need)
        chk.setCursor(Qt.CursorShape.PointingHandCursor)
        chk.setStyleSheet("""
            QCheckBox { border: none; font-weight: bold; color: #333; background-color: transparent; margin-left: 10px; }
            QCheckBox::indicator { width: 40px; height: 20px; }
        """)
        
        row_layout.addLayout(info_vbox, stretch=1)
        row_layout.addWidget(chk)
        scroll_layout.addWidget(frame)
        
        checkboxes.append((e, chk))
        
    scroll_layout.addStretch()
    scroll.setWidget(scroll_widget)
    layout.addWidget(scroll)
    
    btn_confirm = QPushButton(tr_confirm)
    btn_confirm.setMinimumHeight(45)
    btn_confirm.setCursor(Qt.CursorShape.PointingHandCursor)
    btn_confirm.setStyleSheet("""
        QPushButton { font-weight: bold; font-size: 15px; background-color: #0078d7; color: white; border-radius: 5px; }
        QPushButton:hover { background-color: #005a9e; }
    """)
    btn_confirm.clicked.connect(dialog.accept)
    layout.addWidget(btn_confirm)
    
    # 保持窗口在最前
    dialog.setWindowFlags(dialog.windowFlags() | Qt.WindowType.WindowStaysOnTopHint)
    dialog.exec()
    
    updated = False
    for e, chk in checkboxes:
        if chk.isChecked():
            e["mark"] = "No Need Reply"
            updated = True
            
    if updated:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
            ws.delete_rows(2, ws.max_row)
            for e in all_emails:
                ws.append([
                    e["log_time"], 
                    e["mail_time"], 
                    e["direction"], 
                    e["sender"], 
                    e["receiver"], 
                    e["subject"], 
                    e["folder"], 
                    e["reply_time"],
                    e.get("mark", "")
                ])
            wb.save(excel_file)
            print(f"已更新需要忽略的邮件标记。")
        except Exception as ex:
            print(f"更新 Mark 失败: {ex}")

def run_task():
    os.makedirs("data", exist_ok=True)
    os.makedirs("Log", exist_ok=True)
    config_file = "config.json"
    
    if os.path.exists(config_file):
        with open(config_file, "r", encoding="utf-8") as f:
            full_config = json.load(f)
            params = full_config.get("params", {})
    else:
        print("本地不存在 config.json。无法执行。")
        return
            
    target_accounts = []
    target_file = "target_accounts.txt"
    if os.path.exists(target_file):
        try:
            with open(target_file, "r", encoding="utf-8") as f:
                for line in f:
                    acc = line.strip()
                    if acc:
                        target_accounts.append(acc)
        except Exception as e:
            print(f"读取监控邮箱列表失败: {e}")
        
    exclude_emails = []
    exclude_file = "exclude_emails.txt"
    if os.path.exists(exclude_file):
        try:
            with open(exclude_file, "r", encoding="utf-8") as f:
                for line in f:
                    acc = line.strip()
                    if acc:
                        exclude_emails.append(acc)
        except Exception as e:
            print(f"读取排除邮箱列表失败: {e}")
        
    past_hours = params.get('past_hours', 24)
    no_response_threshold = params.get('no_response_threshold', 0)
    max_emails = params.get('max_emails', 500)
    retention_hours = params.get('retention_hours', 720)
    lang = params.get('lang', 'zh')
    
    excel_file = "data/OriginalMail.xlsx"
    log_file = "Log/emails.log"
    widths_file = "data/widths.json"
    
    try:
        import openpyxl
    except ImportError:
        print("未安装 openpyxl 库。请执行 pip install openpyxl")
        return
        
    wb = None
    if not os.path.exists(excel_file):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mails"
        ws.append(["录入时间", "邮件时间", "收发类型", "发件人", "收件人", "邮件标题", "所属文件夹", "Response State", "Mark"])
        
        widths = {'A': 21.5, 'C': 8.875, 'D': 48.25, 'F': 129.75, 'G': 14.0, 'H': 20.0, 'I': 15.0}
        if os.path.exists(widths_file):
            try:
                with open(widths_file, "r", encoding="utf-8") as wf:
                    widths.update(json.load(wf))
            except:
                pass
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        wb.save(excel_file)
        
    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    
    # 确保列名正确
    if ws.cell(row=1, column=8).value != "Response State":
        ws.cell(row=1, column=8).value = "Response State"
    if ws.cell(row=1, column=9).value != "Mark":
        ws.cell(row=1, column=9).value = "Mark"
    
    current_widths = {}
    for col_letter, dim in ws.column_dimensions.items():
        if dim.width:
            current_widths[col_letter] = dim.width
    if current_widths:
        try:
            with open(widths_file, "w", encoding="utf-8") as wf:
                json.dump(current_widths, wf)
        except:
            pass

    # 读取全部现有邮件
    all_emails = []
    seen_keys = set()
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) >= 6:
            t = str(row[1]) if row[1] else ""
            sender = str(row[3]) if row[3] else ""
            receiver = str(row[4]) if row[4] else ""
            subject = str(row[5]) if row[5] else ""
            
            email_dict = {
                "log_time": str(row[0]) if row[0] else "",
                "mail_time": t,
                "direction": str(row[2]) if row[2] else "",
                "sender": sender,
                "receiver": receiver,
                "subject": subject,
                "folder": str(row[6]) if len(row) > 6 and row[6] else "",
                "reply_time": str(row[7]) if len(row) > 7 and row[7] else "",
                "mark": str(row[8]) if len(row) > 8 and row[8] else ""
            }
            key = (receiver, sender, subject, t)
            if key not in seen_keys:
                seen_keys.add(key)
                all_emails.append(email_dict)
    
    with open(log_file, "a", encoding="utf-8") as f:
        try:
            import win32com.client
            outlook = win32com.client.Dispatch("Outlook.Application")
            namespace = outlook.GetNamespace("MAPI")
            
            new_emails_fetched = []
            
            for target_account_name in target_accounts:
                msg = f"\n[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] 开始读取邮箱 [{target_account_name}] 全目录 ..."
                print(msg)
                f.write(msg + "\n")
                
                target_account = None
                for account in namespace.Accounts:
                    if account.DisplayName == target_account_name:
                        target_account = account
                        break
                        
                if not target_account:
                    err = f"  -> 未找到该邮箱账户，跳过。"
                    print(err)
                    f.write(err + "\n")
                    continue
                    
                try:
                    sent_folder = target_account.DeliveryStore.GetDefaultFolder(5)
                    sent_folder_id = sent_folder.EntryID
                except:
                    sent_folder_id = ""
                    
                account_smtp_address = ""
                try:
                    account_smtp_address = target_account.SmtpAddress
                except:
                    pass
                
                try:
                    root_folder = target_account.DeliveryStore.GetRootFolder()
                except:
                    err = "无法获取该账号的根文件夹，跳过。"
                    print(err)
                    f.write(err + "\n")
                    continue
                    
                cutoff_date = datetime.now() - timedelta(hours=past_hours)
                
                fetched_emails, total_count = process_folder(root_folder, cutoff_date, max_emails, sent_folder_id, account_smtp_address, target_account_name, 0)
                
                summary = f"  -> [{target_account_name}] 读取完成，共拉取 {total_count} 封。\n"
                print(summary)
                f.write(summary)
                
                new_emails_fetched.extend(fetched_emails)
                
            # 合并新邮件
            for email in new_emails_fetched:
                key = (email["receiver"], email["sender"], email["subject"], email["mail_time"])
                if key not in seen_keys:
                    seen_keys.add(key)
                    all_emails.append(email)
                    
            # 按时间排序 (旧的在前)
            all_emails.sort(key=lambda x: x["mail_time"])
            
            # 清理超期的邮件
            cutoff_retention_str = (datetime.now() - timedelta(hours=retention_hours)).strftime('%Y-%m-%d %H:%M:%S')
            all_emails = [e for e in all_emails if e["mail_time"] >= cutoff_retention_str]
            
            # 回复状态机处理
            unanswered_customer_emails = [] # list of dicts reference
            
            for e in all_emails:
                sender = e["sender"]
                if e.get("direction") != "发送" and not is_employee(sender, exclude_emails):
                    # 客户发来的邮件
                    e["reply_time"] = "No Response"
                    unanswered_customer_emails.append(e)
                else:
                    # 员工发送的邮件
                    receivers_str = e["receiver"]
                    emp_subject = clean_subject(e["subject"])
                    
                    matched_customer_emails = []
                    
                    for cust_e in unanswered_customer_emails:
                        # 检查员工邮件的收件人中是否包含该客户的发件人
                        cust_sender_cleaned = cust_e["sender"].split("<")[-1].replace(">", "").strip() if "<" in cust_e["sender"] else cust_e["sender"].strip()
                        if not cust_sender_cleaned:
                            cust_sender_cleaned = cust_e["sender"]
                            
                        # 如果是粗略匹配，只要客户的名字或邮箱在收件人字符串里即可
                        if cust_sender_cleaned.lower() in receivers_str.lower():
                            # 进一步检查标题相似度
                            cust_subject = clean_subject(cust_e["subject"])
                            similarity = difflib.SequenceMatcher(None, emp_subject, cust_subject).ratio()
                            if similarity >= 0.8:
                                matched_customer_emails.append(cust_e)
                    
                    if matched_customer_emails:
                        # 匹配成功，这是一封回复邮件，批量核销之前收到的同主题多封邮件
                        for match in matched_customer_emails:
                            match["reply_time"] = e["mail_time"]
                            unanswered_customer_emails.remove(match)
                        e["reply_time"] = "Reply Mail"
                    else:
                        # 未匹配，或者是主动发送的邮件
                        e["reply_time"] = "Normal Mail"
            
            # 根据未回复门限区分 "No Response" 和 "Wait Reply"
            now_dt = datetime.now()
            for e in unanswered_customer_emails:
                try:
                    mail_dt = datetime.strptime(e["mail_time"], '%Y-%m-%d %H:%M:%S')
                    diff_sec = (now_dt - mail_dt).total_seconds()
                    if diff_sec > no_response_threshold:
                        e["reply_time"] = "No Response"
                    else:
                        e["reply_time"] = "Wait Reply"
                except:
                    e["reply_time"] = "No Response"

            # 覆写 Excel
            ws.delete_rows(2, ws.max_row)
            for e in all_emails:
                ws.append([
                    e["log_time"], 
                    e["mail_time"], 
                    e["direction"], 
                    e["sender"], 
                    e["receiver"], 
                    e["subject"], 
                    e["folder"], 
                    e["reply_time"],
                    e.get("mark", "")
                ])
                
            wb.save(excel_file)
            print(f"\\n[完毕] 成功应用回复匹配状态机，共写入 {len(all_emails)} 条记录。")
            f.write(f"\\n[完毕] 成功应用回复匹配状态机，共写入 {len(all_emails)} 条记录。\\n")
            
            # 弹窗提醒处理
            alert_emails = [e for e in all_emails if e.get("reply_time") == "No Response" and str(e.get("mark", "")).strip() != "No Need Reply"]
            if alert_emails:
                show_alert_dialog(alert_emails, all_emails, excel_file, lang)
            
        except Exception as e:
            err = f"读取邮件时发生严重异常:\n{traceback.format_exc()}"
            print(err)
            f.write(err + "\n")
            if wb:
                wb.save(excel_file)

if __name__ == "__main__":
    try:
        run_task()
    except Exception as e:
        print(f"任务执行失败: {e}")
